"""
Temporary isolated module for filling SWOT slides that use placeholder text
instead of tables.

Template placeholders:
  - Placeholder-strengths     -> strengths content
  - Placeholder-weaknesses    -> weaknesses content
  - Placeholder-opportunities -> opportunities content
  - Placeholder-threats       -> threats content

This is a temporary implementation; the feature will be rewritten separately later.
Do not mix with existing table-based fill logic.
"""

import io
from typing import Any

from pptx import Presentation

from shared.logging_config import setup_logging

logger = setup_logging("fill_engine")

# Placeholder strings in the new SWOT template (exact match in shape text)
SWOT_PLACEHOLDERS = {
    "Placeholder-strengths",
    "Placeholder-weaknesses",
    "Placeholder-opportunities",
    "Placeholder-threats",
}


def slide_has_swot_placeholders(slide) -> bool:
    """Check if the slide contains any SWOT placeholder strings."""
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text:
            text = (shape.text or "").strip()
            for placeholder in SWOT_PLACEHOLDERS:
                if placeholder in text:
                    return True
        if hasattr(shape, "has_table") and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    if any(ph in (cell.text or "") for ph in SWOT_PLACEHOLDERS):
                        return True
    return False


def _replace_in_text_frame(text_frame, values: dict[str, str]) -> bool:
    """Replace placeholder strings in a text frame. Returns True if any replacement was made."""
    changed = False
    for para in text_frame.paragraphs:
        for run in para.runs:
            old_text = run.text or ""
            new_text = old_text
            for placeholder, content in values.items():
                if placeholder in new_text:
                    new_text = new_text.replace(placeholder, content)
                    changed = True
            if new_text != old_text:
                run.text = new_text
    return changed


def fill_swot_placeholders(
    pptx_bytes: bytes,
    slide_idx: int,
    strengths: str,
    weaknesses: str,
    opportunities: str,
    threats: str,
) -> bytes:
    """
    Replace SWOT placeholder strings in the slide with actual content.

    Args:
        pptx_bytes: Raw pptx file bytes
        slide_idx: Index of the slide to fill
        strengths, weaknesses, opportunities, threats: Content for each quadrant

    Returns:
        Updated pptx bytes
    """
    values = {
        "Placeholder-strengths": (strengths or "").strip(),
        "Placeholder-weaknesses": (weaknesses or "").strip(),
        "Placeholder-opportunities": (opportunities or "").strip(),
        "Placeholder-threats": (threats or "").strip(),
    }

    prs = Presentation(io.BytesIO(pptx_bytes))
    if slide_idx < 0 or slide_idx >= len(prs.slides):
        raise ValueError(f"Invalid slide_idx: {slide_idx}")

    slide = prs.slides[slide_idx]

    for shape in slide.shapes:
        if hasattr(shape, "has_text_frame") and shape.has_text_frame:
            _replace_in_text_frame(shape.text_frame, values)
        if hasattr(shape, "has_table") and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    _replace_in_text_frame(cell.text_frame, values)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    logger.info(
        "Filled SWOT placeholders on slide %d",
        slide_idx,
        extra={"slide_idx": slide_idx},
    )
    return output.read()


def table_data_to_swot_values(table_data: list[list[Any]]) -> tuple[str, str, str, str]:
    """
    Extract (strengths, weaknesses, opportunities, threats) from table_data.

    SWOT table_data format: [[strengths, weaknesses, opportunities, threats]]
    """
    if not table_data or not isinstance(table_data[0], (list, tuple)):
        return ("", "", "", "")
    row = table_data[0]
    return (
        str(row[0]).strip() if len(row) > 0 else "",
        str(row[1]).strip() if len(row) > 1 else "",
        str(row[2]).strip() if len(row) > 2 else "",
        str(row[3]).strip() if len(row) > 3 else "",
    )
