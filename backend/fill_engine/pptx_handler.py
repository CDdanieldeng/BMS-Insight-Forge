"""PPTX extraction utilities: slide titles and table data as dict."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Optional

from pptx import Presentation
from pptx.table import Table


def load_presentation(path: str | Path | None = None, pptx_bytes: bytes | None = None) -> Presentation:
    """Load a presentation from path or bytes."""
    if pptx_bytes is not None:
        return Presentation(io.BytesIO(pptx_bytes))
    if path is not None:
        return Presentation(str(path))
    raise ValueError("Provide path or pptx_bytes")


def get_slide_count(prs: Presentation) -> int:
    return len(prs.slides)


def get_slide_title(prs: Presentation, slide_idx: int) -> str:
    slide = prs.slides[slide_idx]
    for shape in slide.shapes:
        if shape.has_text_frame and shape.shape_type == 14:  # PLACEHOLDER
            text = shape.text_frame.text.strip()
            if text:
                return text
    return f"Slide {slide_idx + 1}"


def _find_table(prs: Presentation, slide_idx: int) -> Optional[Table]:
    slide = prs.slides[slide_idx]
    for shape in slide.shapes:
        if shape.has_table:
            return shape.table
    return None


def extract_table_as_dict(prs: Presentation, slide_idx: int) -> Optional[dict[str, Any]]:
    """Extract the first table on a slide into a Python-friendly dict.

    Returns None if the slide has no table. Otherwise returns:
    {
        "columns": [...],
        "row_index": [...],
        "data": { row_label: { col_label: cell_text, ... }, ... }
    }
    """
    table = _find_table(prs, slide_idx)
    if table is None:
        return None

    columns = [table.cell(0, c).text.strip() for c in range(1, len(table.columns))]

    row_index = []
    data: dict[str, dict[str, str]] = {}
    for r in range(1, len(table.rows)):
        row_label = table.cell(r, 0).text.strip()
        row_index.append(row_label)
        data[row_label] = {}
        for ci, col_label in enumerate(columns):
            data[row_label][col_label] = table.cell(r, ci + 1).text.strip()

    return {"columns": columns, "row_index": row_index, "data": data}


def fill_table_from_dict(
    prs: Presentation, slide_idx: int, table_data: dict[str, Any]
) -> None:
    """Write data back into PPTX table cells, preserving existing formatting."""
    table = _find_table(prs, slide_idx)
    if table is None:
        return

    columns = table_data["columns"]
    row_index = table_data["row_index"]
    data = table_data["data"]

    for r_idx, row_label in enumerate(row_index):
        table_row = r_idx + 1  # skip header row
        if row_label not in data:
            continue
        for c_idx, col_label in enumerate(columns):
            table_col = c_idx + 1  # skip index column
            cell = table.cell(table_row, table_col)
            new_text = data[row_label].get(col_label, "")
            if not new_text:
                continue
            _set_cell_text_preserve_format(cell, new_text)


def _set_cell_text_preserve_format(cell, text: str) -> None:
    """Set cell text while keeping the first paragraph/run formatting."""
    tf = cell.text_frame
    if tf.paragraphs and tf.paragraphs[0].runs:
        run = tf.paragraphs[0].runs[0]
        run.text = text
        # Remove extra paragraphs if any
        while len(tf.paragraphs) > 1:
            p_elem = tf.paragraphs[-1]._p
            p_elem.getparent().remove(p_elem)
    else:
        para = tf.paragraphs[0] if tf.paragraphs else tf.add_paragraph()
        para.text = text


def save_presentation(prs: Presentation, path: str | Path) -> None:
    prs.save(str(path))
