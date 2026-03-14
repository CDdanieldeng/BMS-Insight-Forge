"""DOCX parser/chunker with heading-path metadata."""

from __future__ import annotations

import tempfile
import uuid
from pathlib import Path

from retriever.models import ChunkRecord
from retriever.parsers.md_parser import parse_markdown_text
from retriever.parsers.normalize import (
    estimate_tokens,
    keep_short_chunk,
    normalize_whitespace,
    p0_clean_lines,
    segment_hints_from_text,
)

_MAX_CHUNK_TOKENS = 500
_MIN_CHUNK_TOKENS = 100


def _split_large(paragraphs: list[str]) -> list[str]:
    chunks: list[str] = []
    cur: list[str] = []
    cur_tokens = 0
    for p in paragraphs:
        p_t = estimate_tokens(p)
        if cur and cur_tokens + p_t > _MAX_CHUNK_TOKENS:
            chunks.append("\n".join(cur))
            cur = [p]
            cur_tokens = p_t
        else:
            cur.append(p)
            cur_tokens += p_t
    if cur:
        chunks.append("\n".join(cur))
    return chunks


def _parse_with_python_docx(content: bytes, *, doc_id: str) -> list[ChunkRecord]:
    from docx import Document

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        temp_path = Path(f.name)
        f.write(content)
        f.flush()
    try:
        doc = Document(str(temp_path))
        chunks: list[ChunkRecord] = []
        heading_stack: list[str] = []
        current_heading = ""
        current_paragraphs: list[str] = []
        para_start = 0

        def flush(para_end: int) -> None:
            nonlocal current_paragraphs, para_start, current_heading
            cleaned = p0_clean_lines(current_paragraphs)
            if not cleaned:
                current_paragraphs = []
                return
            raw_text = normalize_whitespace("\n".join(cleaned))
            pieces = [raw_text] if estimate_tokens(raw_text) <= _MAX_CHUNK_TOKENS else _split_large(cleaned)
            for part_idx, piece in enumerate(pieces):
                token_length = estimate_tokens(piece)
                if token_length < _MIN_CHUNK_TOKENS and not keep_short_chunk(piece, table_flag=False):
                    continue
                chunks.append(
                    ChunkRecord(
                        chunk_id=f"{doc_id}:{uuid.uuid4().hex[:8]}",
                        doc_id=doc_id,
                        source_type="docx",
                        text=piece,
                        token_length=token_length,
                        metadata={
                            "heading_path": list(heading_stack),
                            "heading_title": current_heading or (heading_stack[-1] if heading_stack else ""),
                            "paragraph_range": [para_start, para_end],
                            "part_index": part_idx,
                        },
                        table_flag=False,
                        segment_hint=segment_hints_from_text(piece),
                        noise_flag=False,
                    )
                )
            current_paragraphs = []

        for idx, para in enumerate(doc.paragraphs, start=1):
            text = normalize_whitespace(para.text)
            style_name = (para.style.name if para.style else "").lower()
            if style_name.startswith("heading"):
                flush(idx - 1)
                level_raw = "".join(ch for ch in style_name if ch.isdigit())
                level = int(level_raw) if level_raw else 1
                while len(heading_stack) >= level:
                    heading_stack.pop()
                heading_stack.append(text or f"Heading {level}")
                current_heading = heading_stack[-1]
                para_start = idx + 1
                continue
            if not current_paragraphs:
                para_start = idx
            if text:
                current_paragraphs.append(text)

        flush(len(doc.paragraphs))

        # Append tables as standalone chunks.
        for table_idx, table in enumerate(doc.tables):
            lines: list[str] = []
            for row in table.rows:
                vals = [normalize_whitespace(cell.text) for cell in row.cells]
                vals = [v for v in vals if v]
                if vals:
                    lines.append(" | ".join(vals))
            table_text = normalize_whitespace("\n".join(lines))
            if not table_text:
                continue
            chunks.append(
                ChunkRecord(
                    chunk_id=f"{doc_id}:{uuid.uuid4().hex[:8]}",
                    doc_id=doc_id,
                    source_type="docx",
                    text=table_text,
                    token_length=estimate_tokens(table_text),
                    metadata={
                        "heading_path": list(heading_stack),
                        "heading_title": heading_stack[-1] if heading_stack else "",
                        "table_index": table_idx,
                        "paragraph_range": [0, 0],
                    },
                    table_flag=True,
                    segment_hint=segment_hints_from_text(table_text),
                    noise_flag=False,
                )
            )
        return chunks
    finally:
        temp_path.unlink(missing_ok=True)


def parse_docx_bytes(content: bytes, *, doc_id: str) -> list[ChunkRecord]:
    """
    Parse docx bytes.

    Primary path uses python-docx for heading/paragraph metadata.
    Fallback path uses markdown conversion + markdown parser.
    """
    try:
        return _parse_with_python_docx(content, doc_id=doc_id)
    except Exception:
        from retriever.converter import convert_to_markdown

        md = convert_to_markdown(content, f"{doc_id}.docx")
        return parse_markdown_text(md, doc_id=doc_id, source_type="docx")
