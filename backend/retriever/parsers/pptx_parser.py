"""PPTX parser/chunker with slide-level metadata."""

from __future__ import annotations

import io
import re
import uuid

from pptx import Presentation

from retriever.models import ChunkRecord
from retriever.parsers.normalize import (
    estimate_tokens,
    keep_short_chunk,
    normalize_whitespace,
    p0_clean_lines,
    segment_hints_from_text,
)

_MAX_CHUNK_TOKENS = 500
_MIN_CHUNK_TOKENS = 100
_FOOTER_PATTERNS = [
    re.compile(r"^\s*confidential\b", flags=re.IGNORECASE),
    re.compile(r"^\s*page\s+\d+(\s*/\s*\d+)?\s*$", flags=re.IGNORECASE),
    re.compile(r"^\s*\d+\s*/\s*\d+\s*$"),
]


def _is_noise_text(text: str) -> bool:
    raw = (text or "").strip()
    if not raw:
        return True
    for pattern in _FOOTER_PATTERNS:
        if pattern.search(raw):
            return True
    if len(raw) <= 3 and raw.isdigit():
        return True
    return False


def _extract_slide_title(slide) -> str:
    for shape in slide.shapes:
        if getattr(shape, "is_placeholder", False):
            placeholder = getattr(shape, "placeholder_format", None)
            if placeholder and str(placeholder.type).lower().endswith("title"):
                text = getattr(shape, "text", "").strip()
                if text:
                    return text
    for shape in slide.shapes:
        text = getattr(shape, "text", "").strip()
        if text:
            return text[:120]
    return ""


def _table_to_lines(shape) -> list[str]:
    lines: list[str] = []
    rows = list(shape.table.rows)
    if not rows:
        return lines
    header = [shape.table.cell(0, c).text.strip() for c in range(len(shape.table.columns))]
    lines.append("Table header: " + " | ".join(h for h in header if h))
    for r in range(1, len(rows)):
        row_vals = [shape.table.cell(r, c).text.strip() for c in range(len(shape.table.columns))]
        lines.append("Row: " + " | ".join(v for v in row_vals if v))
    return lines


def _split_text_blocks(blocks: list[str]) -> list[str]:
    chunks: list[str] = []
    current: list[str] = []
    tokens = 0
    for block in blocks:
        t = estimate_tokens(block)
        if current and tokens + t > _MAX_CHUNK_TOKENS:
            chunks.append("\n".join(current))
            current = [block]
            tokens = t
        else:
            current.append(block)
            tokens += t
    if current:
        chunks.append("\n".join(current))
    return chunks


def parse_pptx_bytes(content: bytes, *, doc_id: str) -> list[ChunkRecord]:
    """Parse pptx into slide/body/table chunks."""
    prs = Presentation(io.BytesIO(content))
    chunks: list[ChunkRecord] = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_title = _extract_slide_title(slide)
        body_blocks: list[str] = []
        seen_body: set[str] = set()

        for shape_idx, shape in enumerate(slide.shapes):
            if getattr(shape, "has_table", False):
                table_lines = p0_clean_lines(_table_to_lines(shape))
                if not table_lines:
                    continue
                table_text = normalize_whitespace("\n".join(table_lines))
                if not table_text:
                    continue
                token_length = estimate_tokens(table_text)
                chunks.append(
                    ChunkRecord(
                        chunk_id=f"{doc_id}:{uuid.uuid4().hex[:8]}",
                        doc_id=doc_id,
                        source_type="pptx",
                        text=table_text,
                        token_length=token_length,
                        metadata={
                            "slide_number": slide_idx,
                            "slide_title": slide_title,
                            "shape_index": shape_idx,
                        },
                        table_flag=True,
                        segment_hint=segment_hints_from_text(table_text),
                        noise_flag=False,
                    )
                )
                continue

            text = getattr(shape, "text", "")
            cleaned = normalize_whitespace(text)
            if _is_noise_text(cleaned):
                continue
            key = cleaned.lower()
            if key in seen_body:
                continue
            seen_body.add(key)
            body_blocks.append(cleaned)

        if not body_blocks:
            continue
        for part_idx, part in enumerate(_split_text_blocks(body_blocks)):
            token_length = estimate_tokens(part)
            if token_length < _MIN_CHUNK_TOKENS and not keep_short_chunk(part, table_flag=False):
                continue
            chunks.append(
                ChunkRecord(
                    chunk_id=f"{doc_id}:{uuid.uuid4().hex[:8]}",
                    doc_id=doc_id,
                    source_type="pptx",
                    text=part,
                    token_length=token_length,
                    metadata={
                        "slide_number": slide_idx,
                        "slide_title": slide_title,
                        "part_index": part_idx,
                    },
                    table_flag=False,
                    segment_hint=segment_hints_from_text(part),
                    noise_flag=False,
                )
            )
    return chunks
