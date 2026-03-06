import io
import unittest
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from retriever.parsers.md_parser import parse_markdown_text
from retriever.parsers.pptx_parser import parse_pptx_bytes

try:
    from docx import Document
except Exception:  # pragma: no cover - optional in some envs
    Document = None

from retriever.parsers.docx_parser import parse_docx_bytes


class TestParsersBasic(unittest.TestCase):
    def test_md_heading_and_table(self):
        md = """# Segment Insights
## Preferences
Safe Player prefers WeChat for clinical updates.

| Segment | Monthly volume |
|---|---|
| Safe Player | 34 patients per month |
"""
        chunks = parse_markdown_text(md, doc_id="md1", source_type="pdf_md")
        self.assertTrue(chunks)
        self.assertTrue(any(c.table_flag for c in chunks))
        self.assertTrue(any("heading_path" in c.metadata for c in chunks))

    def test_pptx_table_parsing(self):
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Slide 1"
        textbox = slide.shapes.add_textbox(0, 0, 5000000, 500000)
        textbox.text_frame.text = "Safe Player prefers conference updates."
        table_shape = slide.shapes.add_table(2, 2, 0, 1000000, 5000000, 1000000)
        table = table_shape.table
        table.cell(0, 0).text = "Segment"
        table.cell(0, 1).text = "Volume"
        table.cell(1, 0).text = "Safe Player"
        table.cell(1, 1).text = "34 patients per month"
        stream = io.BytesIO()
        prs.save(stream)
        chunks = parse_pptx_bytes(stream.getvalue(), doc_id="ppt1")
        self.assertTrue(chunks)
        self.assertTrue(any(c.table_flag for c in chunks))
        self.assertTrue(any("slide_number" in c.metadata for c in chunks))

    @unittest.skipIf(Document is None, "python-docx not available")
    def test_docx_heading_parsing(self):
        doc = Document()
        doc.add_heading("Customer Segmentation", level=1)
        doc.add_paragraph("Safe Player uses WeChat and journals.")
        stream = io.BytesIO()
        doc.save(stream)
        chunks = parse_docx_bytes(stream.getvalue(), doc_id="doc1")
        self.assertTrue(chunks)
        self.assertTrue(any("heading_path" in c.metadata for c in chunks))


if __name__ == "__main__":
    unittest.main()
